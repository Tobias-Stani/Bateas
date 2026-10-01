"""Catálogo: leer el Excel del distribuidor (cualquier formato), guardarlo y buscar en él."""
import json
import re
import unicodedata
from itertools import islice

from fastapi import HTTPException
from openpyxl import load_workbook

from ..config import CUSTOM, EXTRA, FIELDS, HEADER_SCAN, IGNORE, KINDS, PAGE_SIZE
from ..db import db


# --- texto ---

def norm(value):
    # "Título / Álbum" -> "titulo album"
    s = unicodedata.normalize("NFD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # 75000.0 -> "75000"; códigos de barras sin notación científica
    return str(value).strip()


def column_letter(i):
    letters = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        letters = chr(65 + r) + letters
    return letters


# --- claves estables ---

KEY_FIELDS = ("artist", "title", "media", "label", "description", "barcode")


def disc_key(d):
    # identifica una edición entre un Excel y el siguiente (el id es la fila y cambia todos los meses).
    # artista + título no alcanza: el mismo disco viene en varias ediciones ("CD + poster", otro código),
    # y el código de barras solo tampoco: los distribuidores lo repiten entre ediciones
    return "v2:" + "|".join(norm(d.get(f)) for f in KEY_FIELDS)


def legacy_key(d):
    # clave de antes de v2 (artista + título + formato): solo para migrar secciones viejas
    return "|".join(norm(d.get(f)) for f in ("artist", "title", "media"))


def custom_key(name):
    # "Precio USD" -> "precio_usd": así se usa como {variable} en el mensaje
    return norm(name).replace(" ", "_")


# --- lectura del Excel ---

def guess(header):
    h = f" {norm(header)} "
    for field, words in FIELDS.items():
        if any(f" {w} " in h for w in words):
            return field
    return None


def open_sheet(path):
    try:
        return load_workbook(path, read_only=True, data_only=True).active
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo. Tiene que ser un Excel .xlsx.")


def detect(path, saved):
    """Encuentra la fila de encabezados y propone un campo para cada columna.

    saved: lo que la disquería eligió otras veces, por nombre de encabezado normalizado; le gana a la adivinanza.
    """
    ws = open_sheet(path)
    top = [[cell(v) for v in row] for row in islice(ws.iter_rows(values_only=True), HEADER_SCAN + 5)]
    ws.parent.close()
    best, best_score = None, 0
    for i, row in enumerate(top[:HEADER_SCAN]):
        score = len({guess(v) for v in row if guess(v)})
        if score > best_score:
            best, best_score = i, score
    if best is None:  # nada reconocible: la primera fila con datos
        best = next((i for i, row in enumerate(top) if any(row)), 0)
    header, samples = top[best] if top else [], top[best + 1:best + 6]
    width = max([len(header), *(len(r) for r in samples)] or [0])
    used, columns = set(), []
    for i in range(width):
        name = header[i] if i < len(header) else ""
        values = [r[i] for r in samples if i < len(r) and r[i]]
        field = saved.get(norm(name)) or guess(name)
        if field in FIELDS and field in used:
            field = None  # cada campo una sola vez: gana la primera columna
        if field not in FIELDS and field not in KINDS:
            field = CUSTOM if name and values else IGNORE
        used.add(field)
        columns.append({"index": i, "letter": column_letter(i), "header": name, "field": field, "samples": values[:3]})
    return {"header_row": best + 1, "columns": columns}


def read_discs(path, header_row, mapping):
    """Filas listas para guardar y las columnas propias de la disquería."""
    ws = open_sheet(path)
    rows = ws.iter_rows(values_only=True)
    header = [cell(v) for v in next(islice(rows, header_row - 1, None), [])]
    name = lambda i: (header[i] if i < len(header) else "") or f"Columna {column_letter(i)}"
    fields = {i: f for i, f in mapping.items() if f in FIELDS}
    extras = [i for i, f in mapping.items() if f == EXTRA]
    customs = [i for i, f in mapping.items() if f == CUSTOM]
    out = []
    # id = número de fila del Excel, así la disquería lo encuentra directo
    for n, row in enumerate(rows, start=header_row + 1):
        value = lambda i: cell(row[i]) if i < len(row) else ""
        rec = {f: value(i) for i, f in fields.items()}
        if not (rec.get("artist") or rec.get("title")):
            continue
        # columnas extra ("Insert: Sí") van a la descripción
        more = [f"{name(i)}: {value(i)}" for i in extras if value(i)]
        rec["description"] = " · ".join(filter(None, [rec.get("description", ""), *more]))
        own = {name(i): value(i) for i in customs if value(i)}
        out.append((n, *(rec.get(f, "") for f in FIELDS), json.dumps(own, ensure_ascii=False) if own else "", disc_key(rec)))
    ws.parent.close()
    columns = [{"key": custom_key(name(i)) or column_letter(i).lower(), "name": name(i)} for i in customs]
    return out, columns


def validate_plan(header_row, mapping):
    chosen = [f for f in mapping.values() if f in FIELDS]
    if any(f not in FIELDS and f not in KINDS for f in mapping.values()):
        raise HTTPException(400, "Hay una columna con un campo desconocido.")
    if len(chosen) != len(set(chosen)):
        raise HTTPException(400, "Hay dos columnas asignadas al mismo campo.")
    if "artist" not in chosen or "title" not in chosen:
        raise HTTPException(400, "Elegí qué columna es el artista y cuál el título.")
    if not 1 <= header_row <= HEADER_SCAN:
        raise HTTPException(400, "La fila de encabezados no es válida.")


# --- guardado ---

def save(t, rows):
    con = db(t)
    with con:
        con.execute("DROP TABLE IF EXISTS discos")
        con.execute(f"CREATE TABLE discos (id INTEGER PRIMARY KEY, {', '.join(FIELDS)}, extra, k)")
        con.executemany(f"INSERT INTO discos VALUES ({', '.join('?' * (len(FIELDS) + 3))})", rows)
        con.execute("CREATE INDEX discos_k ON discos (k)")
    con.close()


def drop(t):
    con = db(t)
    with con:
        con.execute("DROP TABLE IF EXISTS discos")
    con.close()


# --- consultas ---

def exists(t):
    con = db(t)
    ok = con.execute("SELECT 1 FROM sqlite_master WHERE name = 'discos'").fetchone()
    con.close()
    return bool(ok)


def total(t):
    if not exists(t):
        return 0
    con = db(t)
    n = con.execute("SELECT COUNT(*) FROM discos").fetchone()[0]
    con.close()
    return n


def disc(row):
    d = dict(row)
    d.pop("k", None)
    d["extra"] = json.loads(d.get("extra") or "{}")  # catálogos viejos no tienen la columna
    return d


def search(t, q="", media="", genre="", page=1, per_page=PAGE_SIZE):
    per_page = min(max(per_page, 1), PAGE_SIZE)  # la tienda elige de 10 a 50; nunca más que el máximo
    if not exists(t):
        return {"total": 0, "items": [], "page_size": per_page}
    where, params = [], []
    # cada palabra tiene que aparecer en artista, título o sello
    # ponytail: LIKE full scan, ~60k filas va sobrado; FTS5 si el catálogo crece mucho
    for word in q.split():
        # "#34411" o "34411" también encuentra el disco por su número (la fila del Excel)
        if (n := word.lstrip("#")).isdigit():
            where.append("(id = ? OR (artist || ' ' || title || ' ' || label) LIKE ?)")
            params += [int(n), f"%{n}%"]
        else:
            where.append("(artist || ' ' || title || ' ' || label) LIKE ?")
            params.append(f"%{word}%")
    if media:
        where.append("media = ?")
        params.append(media)
    if genre:
        where.append("genre = ?")
        params.append(genre)
    sql_where = f"WHERE {' AND '.join(where)}" if where else ""
    con = db(t)
    count = con.execute(f"SELECT COUNT(*) FROM discos {sql_where}", params).fetchone()[0]
    items = [disc(r) for r in con.execute(f"SELECT * FROM discos {sql_where} ORDER BY id LIMIT ? OFFSET ?",
                                          [*params, per_page, (max(page, 1) - 1) * per_page])]
    con.close()
    return {"total": count, "items": items, "page_size": per_page}


def filters(t):
    if not exists(t):
        return {"media": [], "genre": []}
    con = db(t)
    res = {c: [r[0] for r in con.execute(f"SELECT DISTINCT {c} FROM discos WHERE {c} != '' ORDER BY {c}")]
           for c in ("media", "genre")}
    con.close()
    return res


def stats(t):
    if not exists(t):
        return {"total": 0, "media": [], "genres": 0}
    con = db(t)
    res = {"total": con.execute("SELECT COUNT(*) FROM discos").fetchone()[0],
           "media": [dict(r) for r in con.execute("SELECT media, COUNT(*) n FROM discos GROUP BY media ORDER BY n DESC")],
           "genres": con.execute("SELECT COUNT(DISTINCT genre) FROM discos").fetchone()[0]}
    con.close()
    return res
